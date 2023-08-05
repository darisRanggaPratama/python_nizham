from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Series, Reference

wb = Workbook()
ws = wb.active
treeDatum = [
	["Car", "Color", "Price"],
	["Toyota", "Silver", 500],
	["Honda", "Black", 340],
	["Suzuki", "Blue", 410]
	]

for row in treeDatum:
	ws.append(row)

ft = Font(bold=True)

for row in ws["A1:C1"]:
	for cell in row:
		cell.font = ft

chart = BarChart()
chart.type = "col"
chart.title = "Car Sales"
chart.y_axis.title = 'price'
chart.x_axis.title = 'Mark'
chart.legend = None

datum = Reference(ws, min_col=3, min_row=2, max_row=4, max_col=3)
categories = Reference(ws, min_col=1, min_row=2, max_row=4, max_col=1)

chart.add_data(datum)
chart.set_categories(categories)

ws.add_chart(chart, "F1")
wb.save("TreeDatum.xlsx")




